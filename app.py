from flask import Flask, jsonify, render_template, request
import json
import os
import unicodedata
import openpyxl

app = Flask(
    __name__,
    static_folder="static",
    template_folder="templates"
)

JSON_FILE = os.path.join("data", "escuela.json")
ARCHIVO_EXCEL = "CALIFICACIONES.xlsx"

NOTA_APROBACION = 6


def cargar_datos():
    if not os.path.exists(JSON_FILE):
        return {}
    with open(JSON_FILE, "r", encoding="utf-8") as archivo:
        return json.load(archivo)


def normalizar(texto):
    texto = unicodedata.normalize("NFD", str(texto))
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return texto.lower().strip()


def a_nota(valor):
    if valor is None:
        return None
    try:
        nota = float(str(valor).strip().replace(",", "."))
    except ValueError:
        return None
    if nota != nota or not (1 <= nota <= 10):
        return None
    return nota


def media(notas):
    notas = [n for n in notas if n is not None]
    if not notas:
        return None
    return round(sum(notas) / len(notas), 2)


def parsear_pendientes(texto):
    texto = str(texto or "").strip()
    if texto.lower() in ("", "-", "0", "no", "ninguna", "ninguno"):
        return []
    partes = [p.strip() for p in texto.split(",") if p.strip()]
    if partes and partes[0].isdigit():
        partes = partes[1:]
    return partes


def notas_cuatrimestre(materia):
    return (
        a_nota(materia.get("nota_primer_cuatrimestre")),
        a_nota(materia.get("nota_segundo_cuatrimestre"))
    )


def analizar_alumno(alumno):
    promedios_materias = {}
    materias_desaprobadas = []
    notas_1 = []
    notas_2 = []

    for nombre, materia in alumno.get("materias", {}).items():
        n1, n2 = notas_cuatrimestre(materia)
        if n1 is not None:
            notas_1.append(n1)
        if n2 is not None:
            notas_2.append(n2)
        propias = [n for n in (n1, n2) if n is not None]
        promedios_materias[nombre] = media(propias)
        if any(n < NOTA_APROBACION for n in propias):
            materias_desaprobadas.append(nombre)

    return {
        "nota_aprobacion": NOTA_APROBACION,
        "pendientes": parsear_pendientes(alumno.get("pendientes")),
        "promedio_general": media(notas_1 + notas_2),
        "promedio_primer_cuatrimestre": media(notas_1),
        "promedio_segundo_cuatrimestre": media(notas_2),
        "promedios_materias": promedios_materias,
        "materias_desaprobadas": materias_desaprobadas
    }


def actualizar_celda_excel(curso, id_alumno, materia_nombre, campo, nuevo_valor):
    """Actualiza la celda exacta en el archivo Excel."""
    if not os.path.exists(ARCHIVO_EXCEL):
        print("Error: No se encuentra el archivo Excel.")
        return False

    wb = openpyxl.load_workbook(ARCHIVO_EXCEL)
    if curso not in wb.sheetnames:
        print(f"Error: La hoja '{curso}' no existe en el Excel.")
        return False

    ws = wb[curso]

    offset_campos = {
        "primer_bimestre": 0,
        "segundo_bimestre": 1,
        "nota_primer_cuatrimestre": 2,
        "tercer_bimestre": 4,
        "cuarto_bimestre": 5,
        "nota_segundo_cuatrimestre": 6,
        "primer_cierre": 7,
        "segundo_cierre": 8,
        "tercer_cierre": 9
    }

    if campo not in offset_campos:
        print(f"Error: El campo '{campo}' no es válido.")
        return False

    # 1. Buscar la columna de la materia (Fila 3)
    col_inicio_materia = None
    col = 5
    while col <= ws.max_column:
        val = ws.cell(row=3, column=col).value
        if val and unicodedata.normalize("NFKD", str(val)).strip().lower() == unicodedata.normalize("NFKD", materia_nombre).strip().lower():
            col_inicio_materia = col
            break
        col += 1

    if not col_inicio_materia:
        print(f"Error: No se encontró la materia '{materia_nombre}' en la fila 3 de la hoja '{curso}'.")
        return False

    # 2. Buscar la fila del alumno comparando IDs de forma limpia (removiendo decimales tipo .0)
    id_buscado = str(id_alumno).strip()
    if id_buscado.endswith(".0"):
        id_buscado = id_buscado[:-2]

    fila_alumno = None
    for r in range(5, ws.max_row + 1):
        id_cell = ws.cell(row=r, column=2).value
        if id_cell is not None:
            id_en_excel = str(id_cell).strip()
            if id_en_excel.endswith(".0"):
                id_en_excel = id_en_excel[:-2]
            
            if id_en_excel == id_buscado:
                fila_alumno = r
                break

    if not fila_alumno:
        print(f"Error: No se encontró al alumno con ID '{id_alumno}' en la hoja '{curso}'.")
        return False

    columna_exacta = col_inicio_materia + offset_campos[campo]

    nuevo_valor_str = str(nuevo_valor).strip()
    if nuevo_valor_str == "" or nuevo_valor_str == "–":
        valor_guardar = None
    else:
        try:
            val_num = float(nuevo_valor_str.replace(",", "."))
            valor_guardar = int(val_num) if val_num.is_integer() else val_num
        except ValueError:
            valor_guardar = nuevo_valor_str

    ws.cell(row=fila_alumno, column=columna_exacta).value = valor_guardar
    wb.save(ARCHIVO_EXCEL)
    print(f"¡Éxito! Actualizado -> Fila {fila_alumno}, Columna {columna_exacta} con valor: {valor_guardar}")
    return True


@app.route("/")
def inicio():
    return render_template("index.html")


@app.route("/api/cursos")
def obtener_cursos():
    datos = cargar_datos()
    return jsonify(list(datos.keys()))


@app.route("/api/curso/<curso>")
def obtener_alumnos(curso):
    datos = cargar_datos()
    if curso not in datos:
        return jsonify([])
    alumnos = []
    for alumno in datos[curso]["alumnos"]:
        alumnos.append({
            "id": alumno["id"],
            "nombre": alumno["nombre"],
            "pendientes": len(parsear_pendientes(alumno.get("pendientes")))
        })
    return jsonify(alumnos)


@app.route("/api/curso/<curso>/resumen")
def resumen_curso(curso):
    datos = cargar_datos()
    if curso not in datos:
        return jsonify({"error": "Curso no encontrado"}), 404

    alumnos = datos[curso]["alumnos"]
    notas_por_materia = {}
    notas_1 = []
    notas_2 = []
    con_pendientes = 0
    con_desaprobadas = 0

    for alumno in alumnos:
        analisis = analizar_alumno(alumno)
        if analisis["pendientes"]:
            con_pendientes += 1
        if analisis["materias_desaprobadas"]:
            con_desaprobadas += 1

        for nombre, materia in alumno.get("materias", {}).items():
            n1, n2 = notas_cuatrimestre(materia)
            lista = notas_por_materia.setdefault(nombre, [])
            if n1 is not None:
                lista.append(n1)
                notas_1.append(n1)
            if n2 is not None:
                lista.append(n2)
                notas_2.append(n2)

    materias = [
        {"nombre": nombre, "promedio": media(notas)}
        for nombre, notas in notas_por_materia.items()
    ]
    materias.sort(key=lambda m: (m["promedio"] is None, m["promedio"] or 0))

    return jsonify({
        "curso": curso,
        "nota_aprobacion": NOTA_APROBACION,
        "cantidad_alumnos": len(alumnos),
        "con_pendientes": con_pendientes,
        "con_desaprobadas": con_desaprobadas,
        "promedio_general": media(notas_1 + notas_2),
        "promedio_primer_cuatrimestre": media(notas_1),
        "promedio_segundo_cuatrimestre": media(notas_2),
        "materias": materias
    })


@app.route("/api/alumno/<curso>/<id_alumno>")
def obtener_alumno(curso, id_alumno):
    datos = cargar_datos()
    if curso not in datos:
        return jsonify({"error": "Curso no encontrado"}), 404

    for alumno in datos[curso]["alumnos"]:
        if str(alumno["id"]) == str(id_alumno):
            resultado = dict(alumno)
            resultado["analisis"] = analizar_alumno(alumno)
            return jsonify(resultado)

    return jsonify({"error": "Alumno no encontrado"}), 404


@app.route("/api/alumnos/buscar")
def buscar_alumnos():
    q = normalizar(request.args.get("q", ""))
    if len(q) < 2:
        return jsonify([])

    palabras = q.split()
    datos = cargar_datos()
    resultados = []

    for nombre_curso, curso in datos.items():
        for alumno in curso.get("alumnos", []):
            nombre_normalizado = normalizar(alumno.get("nombre", ""))
            if all(p in nombre_normalizado for p in palabras):
                resultados.append({
                    "id": alumno["id"],
                    "nombre": alumno["nombre"],
                    "curso": nombre_curso,
                    "pendientes": len(parsear_pendientes(alumno.get("pendientes")))
                })

    resultados.sort(key=lambda a: normalizar(a["nombre"]))
    return jsonify(resultados[:100])


@app.route("/api/actualizar", methods=["POST"])
def actualizar_calificacion():
    data = request.json or {}
    curso = data.get("curso")
    id_alumno = str(data.get("id_alumno"))
    materia = data.get("materia")
    campo = data.get("campo")
    valor = str(data.get("valor", "")).strip()

    exito_excel = actualizar_celda_excel(curso, id_alumno, materia, campo, valor)
    if not exito_excel:
        return jsonify({"error": "No se pudo actualizar el archivo Excel."}), 400

    datos = cargar_datos()
    if curso in datos:
        for alumno in datos[curso]["alumnos"]:
            if str(alumno["id"]) == id_alumno:
                if materia in alumno.get("materias", {}):
                    alumno["materias"][materia][campo] = valor if valor != "" else None
                break

    with open(JSON_FILE, "w", encoding="utf-8") as archivo:
        json.dump(datos, archivo, ensure_ascii=False, indent=4, default=str)

    return jsonify({"success": True, "mensaje": "Actualizado correctamente"})


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)